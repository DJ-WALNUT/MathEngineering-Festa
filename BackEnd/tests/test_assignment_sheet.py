"""배정 엑셀 양식 내려받기 · 채워서 올리기.

조 편성은 엑셀에서 끝나고 그 결과가 시스템으로 넘어온다. 이 파일이 지키는 계약은
**왕복**이다 — 내려받은 파일에 값만 채워 그대로 올리면 반영되어야 하고, 채우지 않은
칸이 기존 배정을 지워서는 안 된다.
"""

from __future__ import annotations

import io
import os
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

os.environ.setdefault("SECRET_KEY", "test-secret")
os.environ.setdefault("INGEST_TOKEN", "test-token")
os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")

from app import create_app  # noqa: E402
from app.assignments import ensure_default_fields  # noqa: E402
from app.extensions import db  # noqa: E402
from app.security import hash_password  # noqa: E402

PASSWORD = "mt-admin-1234"
# 최고 관리자 로그인 ID. 비밀번호는 모두 같고 ID 로 사람을 구분한다.
SUPER_ID = "wont0309"
INGEST_HEADERS = {"X-Ingest-Token": "test-token"}


@pytest.fixture()
def client():
    application = create_app()
    application.config.update(
        TESTING=True,
        ADMIN_PASSWORD_HASH=hash_password(PASSWORD),
        # 등급별 비밀번호는 test_admin_tiers.py 가 따로 지킨다. 여기서는 둘을 같게 두어
        # 이 파일이 보는 계약(누가 했는지 · 무엇이 보이는지)만 남긴다.
        STAFF_PASSWORD_HASH=hash_password(PASSWORD),
        LOGIN_RATE_LIMIT="1000 per minute",
        LOOKUP_RATE_LIMIT="1000 per minute",
    )
    with application.app_context():
        db.drop_all()
        db.create_all()
        ensure_default_fields()
        yield application.test_client()
        db.session.remove()


def auth_headers(client) -> dict:
    response = client.post("/api/admin/login", json={"username": SUPER_ID, "password": PASSWORD})
    assert response.status_code == 200, response.get_json()
    return {"Authorization": f"Bearer {response.get_json()['token']}"}


def add_person(client, headers, name: str, phone: str, student_id: str, confirm: bool = True) -> int:
    """수기 등록 → 면제 처리 → 참가 확정. 반환값은 참가자 id."""
    response = client.post(
        "/api/admin/participants",
        headers=headers,
        json={"name": name, "phone": phone, "studentId": student_id},
    )
    assert response.status_code == 201, response.get_json()
    pid = response.get_json()["participant"]["id"]

    if confirm:
        client.patch(
            f"/api/admin/participants/{pid}", headers=headers, json={"statusOverride": "WAIVED"}
        )
        confirmed = client.patch(
            f"/api/admin/participants/{pid}", headers=headers, json={"isConfirmed": True}
        )
        assert confirmed.status_code == 200, confirmed.get_json()
    return pid


def add_options(client, headers, key: str, *values: str) -> None:
    fields = client.get("/api/admin/assignment-fields", headers=headers).get_json()["items"]
    field = next(item for item in fields if item["key"] == key)
    response = client.patch(
        f"/api/admin/assignment-fields/{field['id']}",
        headers=headers,
        json={"options": [*field["options"], *values]},
    )
    assert response.status_code == 200, response.get_json()


def add_choice_field(client, headers, label: str, *options: str) -> str:
    """선택지형 배정 항목을 하나 만든다. 반환값은 값 딕셔너리의 key.

    기본 항목은 조 하나뿐이라, 조 말고 다른 열이 필요한 시험은 여기서 만든다.
    """
    created = client.post(
        "/api/admin/assignment-fields",
        headers=headers,
        json={"label": label, "kind": "choice", "options": list(options)},
    )
    assert created.status_code == 201, created.get_json()
    return created.get_json()["key"]


def download(client, headers) -> bytes:
    response = client.get("/api/admin/assignments/template.xlsx", headers=headers)
    assert response.status_code == 200, response.data[:200]
    return response.data


def upload(client, headers, data: bytes, filename: str = "assignments.xlsx"):
    return client.post(
        "/api/admin/assignments/import",
        headers=headers,
        data={"file": (io.BytesIO(data), filename)},
        content_type="multipart/form-data",
    )


def edit(data: bytes, edits: dict[tuple[int, str], str]) -> bytes:
    """{(행, 열문자): 값} 대로 고쳐 다시 바이트로. 사람이 엑셀에서 채우는 것과 같다."""
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook["배정"]
    for (row, column), value in edits.items():
        sheet[f"{column}{row}"] = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def row_of(data: bytes, name: str) -> int:
    """양식에서 그 사람의 행 번호. 명단은 이름순이라 등록 순서와 다르다."""
    sheet = load_workbook(io.BytesIO(data))["배정"]
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row[1] == name:
            return index
    raise AssertionError(f"{name} 을 양식에서 찾지 못했습니다")


def assignments_of(client, headers, pid: int) -> dict:
    return client.get(f"/api/admin/participants/{pid}", headers=headers).get_json()["assignments"]


def test_template_carries_current_fields_and_confirmed_people(client):
    headers = auth_headers(client)
    add_options(client, headers, "group", "1조", "2조")
    add_person(client, headers, "홍길동", "010-1234-5678", "20261111")
    add_person(client, headers, "미확정", "010-2222-3333", "20262222", confirm=False)

    # 운영 중에 늘린 항목도 양식에 함께 나와야 한다
    client.post(
        "/api/admin/assignment-fields",
        headers=headers,
        json={"label": "텐트", "kind": "choice", "options": ["A동", "B동"]},
    )

    workbook = load_workbook(io.BytesIO(download(client, headers)))
    sheet = workbook["배정"]
    header = [cell.value for cell in sheet[1]]
    assert header[:4] == ["참가자ID", "이름", "학번", "학과(참고)"]
    assert "조" in header and "텐트" in header

    names = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert names == ["홍길동"]  # 확정 전인 사람은 배정 대상이 아니다

    # 선택지는 숨긴 시트에 실어 드롭다운으로 걸어 둔다
    assert workbook["선택지"].sheet_state == "hidden"
    assert "사용법" in workbook.sheetnames
    assert len(sheet.data_validations.dataValidation) == 2


def test_round_trip_fills_assignments(client):
    headers = auth_headers(client)
    add_options(client, headers, "group", "1조", "2조")
    seat = add_choice_field(client, headers, "자리", "A구역")
    first = add_person(client, headers, "홍길동", "010-1234-5678", "20261111")
    second = add_person(client, headers, "이순신", "010-2222-3333", "20262222")

    data = download(client, headers)
    # 열 순서: A 참가자ID · B 이름 · C 학번 · D 학과 · E 조 · F 자리
    hong, lee = row_of(data, "홍길동"), row_of(data, "이순신")
    filled = edit(data, {(hong, "E"): "1조", (hong, "F"): "A구역", (lee, "E"): "2조"})

    result = upload(client, headers, filled)
    assert result.status_code == 200, result.get_json()
    body = result.get_json()
    assert body["updated"] == 2
    assert body["changedValues"] == 3
    assert body["skippedCount"] == 0

    assert assignments_of(client, headers, first) == {"group": "1조", seat: "A구역"}
    assert assignments_of(client, headers, second) == {"group": "2조"}

    audit = client.get("/api/admin/audit", headers=headers).get_json()["items"]
    entry = next(item for item in audit if item["action"] == "assignment.import")
    assert any(change["label"] == "홍길동 · 조" for change in entry["detail"]["changes"])


def test_blank_cells_leave_existing_assignments_alone(client):
    """반쯤 채운 파일을 올렸다고 나머지 배정이 지워지면 안 된다."""
    headers = auth_headers(client)
    add_options(client, headers, "group", "1조")
    seat = add_choice_field(client, headers, "자리", "A구역")
    pid = add_person(client, headers, "홍길동", "010-1234-5678", "20261111")

    upload(client, headers, edit(download(client, headers), {(2, "E"): "1조", (2, "F"): "A구역"}))
    assert assignments_of(client, headers, pid) == {"group": "1조", seat: "A구역"}

    # 조만 적고 자리 칸은 비운 파일 — 자리 배정은 그대로여야 한다
    blank = edit(download(client, headers), {(2, "F"): None})
    body = upload(client, headers, blank).get_json()
    assert body["updated"] == 0
    assert body["unchanged"] == 1
    assert assignments_of(client, headers, pid) == {"group": "1조", seat: "A구역"}

    # 풀려면 '-' 라고 적는다
    cleared = edit(download(client, headers), {(2, "F"): "-"})
    assert upload(client, headers, cleared).get_json()["updated"] == 1
    assert assignments_of(client, headers, pid) == {"group": "1조"}


def test_bad_value_skips_only_that_row(client):
    headers = auth_headers(client)
    add_options(client, headers, "group", "1조")
    good = add_person(client, headers, "홍길동", "010-1234-5678", "20261111")
    bad = add_person(client, headers, "이순신", "010-2222-3333", "20262222")

    data = download(client, headers)
    good_row, bad_row = row_of(data, "홍길동"), row_of(data, "이순신")
    filled = edit(data, {(good_row, "E"): "1조", (bad_row, "E"): "9조"})
    body = upload(client, headers, filled).get_json()

    assert body["updated"] == 1
    assert body["skippedCount"] == 1
    assert "9조" in body["skipped"][0]["reason"]
    assert body["skipped"][0]["row"] == bad_row

    assert assignments_of(client, headers, good) == {"group": "1조"}
    assert assignments_of(client, headers, bad) == {}


def test_unconfirmed_person_is_reported_not_assigned(client):
    """확정 전인 사람을 줄에 끼워 올려도 조에 들어가서는 안 된다."""
    headers = auth_headers(client)
    add_options(client, headers, "group", "1조")
    add_person(client, headers, "홍길동", "010-1234-5678", "20261111")
    pending = add_person(client, headers, "미확정", "010-2222-3333", "20262222", confirm=False)

    data = download(client, headers)
    filled = edit(data, {(2, "E"): "1조", (3, "A"): pending, (3, "B"): "미확정", (3, "E"): "1조"})

    body = upload(client, headers, filled).get_json()
    assert body["updated"] == 1
    assert body["skippedCount"] == 1
    assert "확정" in body["skipped"][0]["reason"]
    assert assignments_of(client, headers, pending) == {}


def test_row_without_id_is_matched_by_student_id(client):
    """양식을 다시 만들어 붙여 넣느라 ID 열이 사라지는 일이 잦다."""
    headers = auth_headers(client)
    add_options(client, headers, "group", "1조")
    pid = add_person(client, headers, "홍길동", "010-1234-5678", "20261111")

    filled = edit(download(client, headers), {(2, "A"): None, (2, "E"): "1조"})
    body = upload(client, headers, filled).get_json()

    assert body["updated"] == 1
    assert assignments_of(client, headers, pid) == {"group": "1조"}


def test_ambiguous_name_without_student_id_is_skipped(client):
    """동명이인은 고르지 않고 건너뛴다. 엉뚱한 사람을 조에 넣는 것보다 낫다."""
    headers = auth_headers(client)
    add_options(client, headers, "group", "1조")
    first = add_person(client, headers, "홍길동", "010-1234-5678", "20261111")
    second = add_person(client, headers, "홍길동", "010-2222-3333", "20262222")

    filled = edit(
        download(client, headers),
        {(2, "A"): None, (2, "C"): None, (2, "E"): "1조"},
    )
    body = upload(client, headers, filled).get_json()

    assert body["updated"] == 0
    assert body["skippedCount"] == 1
    assert "여럿" in body["skipped"][0]["reason"]
    assert assignments_of(client, headers, first) == {}
    assert assignments_of(client, headers, second) == {}


def test_sheet_without_recognisable_header_is_rejected(client):
    headers = auth_headers(client)
    add_person(client, headers, "홍길동", "010-1234-5678", "20261111")

    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["아무", "관계없는", "표"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = upload(client, headers, buffer.getvalue(), filename="엉뚱한파일.xlsx")
    assert response.status_code == 400
    assert response.get_json()["error"] == "parse_failed"


def test_csv_is_accepted_too(client):
    """엑셀에서 'CSV로 저장'을 눌러 올리는 사람이 반드시 나온다."""
    headers = auth_headers(client)
    add_options(client, headers, "group", "1조")
    pid = add_person(client, headers, "홍길동", "010-1234-5678", "20261111")

    csv_text = f"참가자ID,이름,학번,조\n{pid},홍길동,20261111,1조\n"
    response = upload(client, headers, csv_text.encode("utf-8-sig"), filename="배정.csv")

    assert response.status_code == 200, response.get_json()
    assert response.get_json()["updated"] == 1
    assert assignments_of(client, headers, pid) == {"group": "1조"}

"""카카오뱅크 거래내역 파일(xlsx/csv) 파서.

카카오뱅크 내려받기 파일은 앞쪽에 계좌 요약 몇 줄이 붙고, 컬럼명도 버전마다
조금씩 다르다. 그래서 헤더 위치를 찾아낸 뒤 별칭 표로 컬럼을 매핑한다.
입금 건만 취하고 출금은 버린다.

내려받은 파일에는 열기 비밀번호가 걸려 있다. 비밀번호를 알려 주면 여기서
직접 풀어 읽으므로, 올리는 사람이 매번 손으로 해제할 필요는 없다.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import datetime

import msoffcrypto
from openpyxl import load_workbook

from ..utils import normalize_name, parse_amount, parse_datetime, split_depositor

# 헤더 별칭 (공백 제거 후 비교)
_ALIASES: dict[str, tuple[str, ...]] = {
    "datetime": ("거래일시", "거래일자", "거래시각", "일시", "날짜", "거래날짜", "transactiondate", "date"),
    "kind": ("구분", "거래구분", "입출금구분", "거래종류", "유형", "거래유형", "type"),
    "amount": ("거래금액", "금액", "amount"),
    "amount_in": ("입금액", "입금금액", "맡기신금액", "입금", "deposit"),
    "amount_out": ("출금액", "출금금액", "찾으신금액", "출금", "withdrawal"),
    "balance": ("거래후잔액", "거래후금액", "잔액", "잔고", "남은잔액", "balance"),
    "name": (
        "내용", "적요", "거래내용", "입금자명", "보낸분", "받는분", "상대방", "거래상대",
        "상대계좌예금주", "메모", "description", "note",
    ),
}

_DEPOSIT_KINDS = ("입금", "입", "받", "deposit", "in")
_HEADER_SCAN_ROWS = 40

# 비밀번호가 걸린 엑셀은 zip(xlsx)이 아니라 OLE 복합문서 껍데기에 담겨 나온다.
# 확장자는 그대로 .xlsx 이므로, 내용을 열기 전에 앞 8바이트로 구분한다.
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# 머리말의 '조회기간  2026.05.01 - 2026.05.26' 에서 조회 범위를 뽑는다.
# 이 범위 밖의 입금은 파일에 아예 없으므로, 어디까지 반영됐는지 알리는 데 쓴다.
_PERIOD_RE = re.compile(
    r"(\d{4}[.\-/]\s?\d{1,2}[.\-/]\s?\d{1,2})\s*[-~]\s*(\d{4}[.\-/]\s?\d{1,2}[.\-/]\s?\d{1,2})"
)


@dataclass
class ParsedRow:
    row_number: int
    occurred_at: datetime
    amount: int
    balance_after: int | None
    raw_name: str
    name_norm: str
    digit_suffix: str
    raw: dict


@dataclass
class StatementResult:
    rows: list[ParsedRow] = field(default_factory=list)
    skipped: list[dict] = field(default_factory=list)
    header_row: int | None = None
    columns: dict[str, int] = field(default_factory=dict)
    # 파일 머리말에 적힌 조회 범위 (없으면 None)
    period_from: datetime | None = None
    period_to: datetime | None = None
    # 비밀번호가 걸린 파일을 서버가 풀어서 읽었는지. 화면에 알려 주는 용도.
    decrypted: bool = False

    @property
    def latest_transaction_at(self) -> datetime | None:
        return max((row.occurred_at for row in self.rows), default=None)


def _norm_header(value) -> str:
    if value is None:
        return ""
    return re.sub(r"[\s_\-()\[\]]", "", str(value)).lower()


def _detect_columns(row: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _norm_header(cell)
        if not key:
            continue
        for field_name, aliases in _ALIASES.items():
            if field_name in mapping:
                continue
            if key in aliases:
                mapping[field_name] = index
                break
    # '내용' 같은 이름 컬럼이 여러 개면 첫 번째만 쓰므로 별도 처리 불필요
    return mapping


def _is_usable_header(mapping: dict[str, int]) -> bool:
    has_time = "datetime" in mapping
    has_money = "amount" in mapping or "amount_in" in mapping
    return has_time and has_money


def _unlock_xlsx(data: bytes, passwords: Sequence[str]) -> bytes:
    """열기 비밀번호를 풀어 알맹이 xlsx 바이트를 돌려준다.

    비밀번호를 여러 개 받는 것은, 계좌를 옮기거나 비밀번호를 바꾼 뒤에도
    옛 파일을 다시 올릴 수 있어야 하기 때문이다. 앞에서부터 차례로 시도한다.
    """
    if not passwords:
        raise ValueError(
            "비밀번호가 걸린 엑셀 파일입니다. 서버에 비밀번호가 등록되어 있지 않아 열지 못했습니다. "
            "(관리자: 백엔드 .env 의 STATEMENT_PASSWORDS 를 채우고 다시 시작해 주세요)"
        )

    for password in passwords:
        buffer = io.BytesIO()
        try:
            office = msoffcrypto.OfficeFile(io.BytesIO(data))
            office.load_key(password=password)
            office.decrypt(buffer)
        except Exception:
            # 비밀번호가 틀리면 InvalidKeyError 지만, 껍데기가 예상과 다르면
            # 파싱·구조체 오류가 그대로 올라온다. 어느 쪽이든 '이 비밀번호로는
            # 못 열었다'는 뜻이라 다음 후보로 넘어간다.
            # 한 번 쓴 OfficeFile 은 상태가 남으므로 매번 새로 만든다.
            continue
        return buffer.getvalue()

    raise ValueError(
        "등록된 비밀번호로 파일을 열지 못했습니다. 파일 비밀번호가 바뀌었는지 확인해 주세요."
    )


# 표를 읽어 오는 두 함수와 OLE 판별은 거래내역 전용이 아니다.
# 배정 엑셀(app/assignment_sheet.py)도 같은 방식으로 파일을 연다.
def read_xlsx_table(stream: io.BytesIO) -> list[list]:
    try:
        workbook = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:
        # 손상된 파일이나 확장자만 xlsx 인 파일. 그대로 두면 500 이 되어
        # 올린 사람은 무슨 일인지 알 수 없다.
        raise ValueError(
            "엑셀 파일을 열지 못했습니다. 파일이 손상되었거나 형식이 다를 수 있습니다."
        ) from exc
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_csv_table(data: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp949", "euc-kr", "utf-8"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    raise ValueError("CSV 인코딩을 인식하지 못했습니다 (utf-8 / cp949 를 지원합니다)")


def _extract_period(preamble: list[list]) -> tuple[datetime | None, datetime | None]:
    """머리말에서 조회 시작·종료일을 찾는다. 실패해도 조용히 넘어간다."""
    for row in preamble:
        for cell in row:
            if cell is None:
                continue
            match = _PERIOD_RE.search(str(cell))
            if not match:
                continue
            start = parse_datetime(match.group(1).replace(" ", ""))
            end = parse_datetime(match.group(2).replace(" ", ""))
            if start or end:
                return start, end
    return None, None


def _cell(row: list, index: int | None):
    if index is None or index >= len(row):
        return None
    return row[index]


def _row_amount(row: list, columns: dict[str, int]) -> tuple[int | None, str]:
    """(입금액, 판정사유). 입금이 아니면 (None, 사유)."""
    amount_in = parse_amount(_cell(row, columns.get("amount_in")))
    if "amount_in" in columns:
        if amount_in and amount_in > 0:
            return amount_in, ""
        amount_out = parse_amount(_cell(row, columns.get("amount_out")))
        if amount_out and amount_out > 0:
            return None, "출금 건"
        if not amount_in:
            return None, "금액 없음"

    kind = str(_cell(row, columns.get("kind")) or "").strip()
    amount = parse_amount(_cell(row, columns.get("amount")))
    if amount is None:
        return None, "금액 없음"
    if kind:
        normalized_kind = _norm_header(kind)
        if not any(token in normalized_kind for token in _DEPOSIT_KINDS):
            return None, f"입금 아님 ({kind})"
        return abs(amount), ""
    # 구분 컬럼이 없으면 부호로 판단한다
    if amount <= 0:
        return None, "출금 건(음수 금액)"
    return amount, ""


def parse_statement(
    data: bytes, filename: str, passwords: Sequence[str] = ()
) -> StatementResult:
    lowered = (filename or "").lower()
    decrypted = False
    if lowered.endswith((".xlsx", ".xlsm")):
        if data[:8] == OLE_MAGIC:
            data = _unlock_xlsx(data, passwords)
            decrypted = True
        table = read_xlsx_table(io.BytesIO(data))
    elif lowered.endswith((".csv", ".txt")):
        table = read_csv_table(data)
    elif lowered.endswith(".xls"):
        raise ValueError("구형 .xls 는 지원하지 않습니다. .xlsx 또는 .csv 로 내려받아 주세요.")
    else:
        raise ValueError("지원하지 않는 파일 형식입니다 (.xlsx / .csv)")

    result = StatementResult(decrypted=decrypted)

    for index, row in enumerate(table[:_HEADER_SCAN_ROWS]):
        mapping = _detect_columns(list(row))
        if _is_usable_header(mapping):
            result.header_row = index
            result.columns = mapping
            break

    if result.header_row is None:
        raise ValueError(
            "거래내역 헤더를 찾지 못했습니다. '거래일시'와 '거래금액'(또는 '입금액') 컬럼이 있는지 확인해 주세요."
        )

    result.period_from, result.period_to = _extract_period(table[: result.header_row])

    columns = result.columns
    header = [str(cell) if cell is not None else "" for cell in table[result.header_row]]

    for offset, raw_row in enumerate(table[result.header_row + 1 :], start=result.header_row + 2):
        row = list(raw_row)
        if not any(cell not in (None, "") for cell in row):
            continue

        occurred_at = parse_datetime(_cell(row, columns.get("datetime")))
        if occurred_at is None:
            result.skipped.append({"row": offset, "reason": "거래일시를 읽지 못했습니다"})
            continue

        amount, reason = _row_amount(row, columns)
        if amount is None:
            result.skipped.append({"row": offset, "reason": reason})
            continue

        balance = parse_amount(_cell(row, columns.get("balance")))
        raw_name = str(_cell(row, columns.get("name")) or "").strip()
        name_part, digits = split_depositor(raw_name)

        result.rows.append(
            ParsedRow(
                row_number=offset,
                occurred_at=occurred_at,
                amount=amount,
                balance_after=balance,
                raw_name=raw_name,
                name_norm=normalize_name(name_part),
                digit_suffix=digits,
                raw={
                    key: (value.isoformat() if isinstance(value, datetime) else value)
                    for key, value in zip(header, row)
                    if value not in (None, "")
                },
            )
        )

    return result
